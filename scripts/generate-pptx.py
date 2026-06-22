#!/usr/bin/env python3
"""
generate-pptx.py — Build gcc-health-ai-profile.pptx from the HTML deck content.
Run: python3 scripts/generate-pptx.py
Output: gcc-health-ai-profile.pptx (in /home/user/slides/)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree
import copy

# ── Colours ────────────────────────────────────────
BG_CREAM   = RGBColor(0xf4, 0xef, 0xe6)
BG_DARK    = RGBColor(0x07, 0x1c, 0x11)
BG_DARK2   = RGBColor(0x0a, 0x5c, 0x3a)
ACCENT_G   = RGBColor(0x0e, 0x7a, 0x4e)   # GCC green
ACCENT_GOLD= RGBColor(0xc9, 0xa9, 0x6e)   # gold
ACCENT_GOLDF=RGBColor(0xdf, 0xc1, 0x8a)   # gold light
FG_DARK    = RGBColor(0x1e, 0x2d, 0x4a)
FG_MUTED   = RGBColor(0x6b, 0x72, 0x80)
WHITE      = RGBColor(0xff, 0xff, 0xff)
WHITE70    = RGBColor(0xb3, 0xb3, 0xb3)    # ~70% white approximation
SURFACE    = RGBColor(0xff, 0xff, 0xff)
BORDER     = RGBColor(0xe8, 0xe1, 0xd2)

# ── Slide dimensions — 16:9 widescreen ────────────
W = Inches(13.333)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]   # blank layout


# ── Helper: add text box ───────────────────────────
def txbox(slide, text, left, top, width, height,
          size=18, bold=False, italic=False,
          color=FG_DARK, align=PP_ALIGN.LEFT,
          font="Calibri", wrap=True, word_wrap=True):
    tf = slide.shapes.add_textbox(left, top, width, height)
    tf.word_wrap = word_wrap
    frame = tf.text_frame
    frame.word_wrap = wrap
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return tf


def txbox2(slide, left, top, width, height, wrap=True):
    """Return text frame for manual paragraph construction."""
    tf = slide.shapes.add_textbox(left, top, width, height)
    tf.word_wrap = wrap
    tf.text_frame.word_wrap = wrap
    return tf.text_frame


def para(tf, text, size=14, bold=False, italic=False, color=FG_DARK,
         align=PP_ALIGN.LEFT, font="Calibri", space_before=0):
    from pptx.util import Pt
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.italic = italic
    return p


def fill_solid(shape, rgb):
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def rect(slide, left, top, width, height, rgb, radius=0):
    from pptx.util import Emu
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE = 1; auto_shape_type ROUNDED_RECTANGLE=5
        left, top, width, height
    )
    fill_solid(shape, rgb)
    shape.line.fill.background()   # no border
    return shape


def rounded_rect(slide, left, top, width, height, rgb, line_rgb=None, line_width=Pt(0.75)):
    shape = slide.shapes.add_shape(5, left, top, width, height)  # 5 = rounded rectangle
    fill_solid(shape, rgb)
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def label_text(slide, text, left, top, width=Inches(4), color=ACCENT_GOLD, size=9):
    tf = slide.shapes.add_textbox(left, top, width, Pt(20))
    p = tf.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text.upper()
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = "Calibri"


def heading(slide, text, left, top, width, height,
            size=36, color=WHITE, font="Georgia"):
    tf = txbox2(slide, left, top, width, height)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = font
    return tf


def body_text(slide, text, left, top, width, height,
              size=13, color=WHITE70):
    tf = slide.shapes.add_textbox(left, top, width, height)
    tf.word_wrap = True
    tf.text_frame.word_wrap = True
    p = tf.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return tf


# ══════════════════════════════════════════════════════
#  S1 — HERO COVER
# ══════════════════════════════════════════════════════
def slide_hero(prs):
    sl = prs.slides.add_slide(BLANK)

    # Background — deep green gradient approximated with solid
    bg = rect(sl, 0, 0, W, H, BG_DARK)

    # Green accent overlay strip (right side glow)
    acc = rect(sl, Inches(7), 0, Inches(6.333), H, RGBColor(0x0a, 0x5c, 0x3a))
    acc.fill.solid()
    acc.fill.fore_color.rgb = RGBColor(0x0a, 0x5c, 0x3a)

    # Decorative large "2025" watermark
    txbox(sl, "2025",
          left=Inches(7.5), top=Inches(2.8),
          width=Inches(5.8), height=Inches(5),
          size=180, bold=True, color=RGBColor(0x14, 0x2b, 0x1e),
          font="Calibri")

    # Logo area top-left
    txbox(sl, "GCC  Health AI",
          left=Inches(0.45), top=Inches(0.3),
          width=Inches(3), height=Inches(0.5),
          size=13, bold=True, color=WHITE, font="Calibri")

    # Kicker
    txbox(sl, "— COMMUNITY PROFILE 2025",
          left=Inches(0.45), top=Inches(4.2),
          width=Inches(8), height=Inches(0.4),
          size=10, bold=True, color=ACCENT_GOLDF, font="Calibri")

    # H1
    heading(sl, "The Community\nPowering Healthcare\nin the Gulf",
            left=Inches(0.45), top=Inches(4.55),
            width=Inches(8.5), height=Inches(2.4),
            size=42, color=WHITE, font="Georgia")

    # Subtitle
    body_text(sl,
        "The Gulf's leading network where clinicians, AI builders, founders, "
        "and investors connect — and turn ideas into adoption.",
        left=Inches(0.45), top=Inches(6.85),
        width=Inches(7.5), height=Inches(0.7),
        size=13, color=RGBColor(0xb8, 0xb8, 0xb8))

    # Brand bottom right
    txbox(sl, "GCC Health AI · gcchealth.ai",
          left=Inches(10), top=Inches(7.05),
          width=Inches(3.1), height=Inches(0.35),
          size=9, color=RGBColor(0x66, 0x77, 0x66),
          align=PP_ALIGN.RIGHT, font="Calibri")


# ══════════════════════════════════════════════════════
#  S2 — ABOUT (split layout)
# ══════════════════════════════════════════════════════
def slide_about(prs):
    sl = prs.slides.add_slide(BLANK)

    # Left panel (dark green)
    LEFT_W = Inches(5.6)
    rect(sl, 0, 0, LEFT_W, H, BG_DARK)

    # Eyebrow
    label_text(sl, "Who We Are", Inches(0.45), Inches(0.5), color=ACCENT_GOLDF, size=9)

    # H2
    heading(sl, "GCC\nHealth AI",
            Inches(0.45), Inches(1.1), LEFT_W - Inches(0.9), Inches(2),
            size=46, font="Georgia")

    # Divider
    rect(sl, Inches(0.45), Inches(3.0), Inches(0.65), Inches(0.06), ACCENT_GOLD)

    # Big stat
    heading(sl, "300+",
            Inches(0.45), Inches(3.2), Inches(3), Inches(1.2),
            size=72, font="Georgia")
    txbox(sl, "Members",
          Inches(0.45), Inches(4.25), Inches(2.5), Inches(0.5),
          size=20, color=ACCENT_GOLDF, bold=True, font="Georgia")
    txbox(sl, "ACROSS 6 GLOBAL MARKETS",
          Inches(0.45), Inches(4.75), Inches(4), Inches(0.35),
          size=9, color=RGBColor(0x88, 0x88, 0x88), bold=True, font="Calibri")
    txbox(sl, "EST. 2023",
          Inches(0.45), Inches(7.1), Inches(2), Inches(0.3),
          size=9, color=RGBColor(0x55, 0x66, 0x55), bold=True, font="Calibri")

    # Right panel
    right_x = LEFT_W + Inches(0.05)
    right_w = W - right_x
    rect(sl, right_x, 0, right_w, H, BG_CREAM)

    label_text(sl, "What We Do", right_x + Inches(0.4), Inches(0.55), color=ACCENT_G, size=9)

    bullets = [
        ("The region's definitive network",
         "One room for clinicians, AI builders, founders, and investors shaping Gulf healthcare"),
        ("Cross-border by design",
         "Partnerships, deal flow, and knowledge moving freely across six markets that usually work alone"),
        ("Built for adoption, not hype",
         "We help AI move from pilot to patient — aligned with Vision 2030's healthcare ambitions"),
        ("Curated, never crowded",
         "Every member is vetted, so conversations stay high-signal and connections actually convert"),
    ]

    y = Inches(1.05)
    for i, (title, desc) in enumerate(bullets):
        # Numbered circle (approximated with a rounded rect)
        circ = rounded_rect(sl, right_x + Inches(0.4), y + Inches(0.04),
                             Inches(0.36), Inches(0.36), ACCENT_G)
        txbox(sl, str(i + 1),
              right_x + Inches(0.43), y,
              Inches(0.3), Inches(0.45),
              size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        txbox(sl, title,
              right_x + Inches(0.9), y - Inches(0.02),
              right_w - Inches(1.2), Inches(0.35),
              size=13, bold=True, color=FG_DARK, font="Calibri")
        txbox(sl, desc,
              right_x + Inches(0.9), y + Inches(0.3),
              right_w - Inches(1.2), Inches(0.5),
              size=11, color=FG_MUTED, font="Calibri")
        y += Inches(1.25)


# ══════════════════════════════════════════════════════
#  S3 — PROBLEM (dark, 4 cards)
# ══════════════════════════════════════════════════════
def slide_problem(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_DARK)

    label_text(sl, "The Challenge", Inches(0.5), Inches(0.45), color=ACCENT_GOLDF, size=9)

    heading(sl, "GCC healthcare is fragmented\n— and AI won't wait",
            Inches(0.5), Inches(0.75), Inches(12), Inches(1.6),
            size=32, font="Georgia")

    problems = [
        ("1", "No Discovery Layer",
         "Founders can't find the right hospital, regulator, or investor — so the right people never meet"),
        ("2", "Markets in Silos",
         "Saudi, the UAE, and Oman each move alone — duplicating effort instead of compounding it"),
        ("3", "Pilots That Stall",
         "AI advances faster than legacy systems adopt it — promising tools die after the pilot"),
        ("4", "Capital Misses Talent",
         "Investors overlook regional deals while founders chase capital abroad"),
    ]

    col_w = Inches(6.0)
    row_h = Inches(2.2)
    xs = [Inches(0.5), Inches(6.9)]
    ys = [Inches(2.4), Inches(4.7)]

    for idx, (num, title, desc) in enumerate(problems):
        col = idx % 2
        row = idx // 2
        x = xs[col]
        y = ys[row]

        card = rounded_rect(sl, x, y, col_w, row_h,
                            RGBColor(0x0d, 0x2a, 0x1b),
                            line_rgb=RGBColor(0x1e, 0x4a, 0x33))
        # Gold top bar
        rect(sl, x, y, col_w, Inches(0.06), ACCENT_GOLD)

        # Number circle
        circ = rounded_rect(sl, x + Inches(0.22), y + Inches(0.35),
                            Inches(0.42), Inches(0.42), ACCENT_G)
        txbox(sl, num, x + Inches(0.27), y + Inches(0.3),
              Inches(0.32), Inches(0.45), size=13, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER)

        txbox(sl, title, x + Inches(0.78), y + Inches(0.3),
              col_w - Inches(1.0), Inches(0.45),
              size=14, bold=True, color=WHITE)
        txbox(sl, desc, x + Inches(0.78), y + Inches(0.72),
              col_w - Inches(1.0), Inches(1.3),
              size=11, color=RGBColor(0x8a, 0xa0, 0x90))


# ══════════════════════════════════════════════════════
#  S4 — MEMBERS
# ══════════════════════════════════════════════════════
def slide_members(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_CREAM)

    label_text(sl, "The Network", Inches(0.5), Inches(0.45), color=ACCENT_GOLD, size=9)
    heading(sl, "Who's Inside",
            Inches(0.5), Inches(0.7), Inches(8), Inches(1.0),
            size=38, color=FG_DARK, font="Georgia")

    # Stat column left
    heading(sl, "300+",
            Inches(0.5), Inches(1.85), Inches(2.5), Inches(1.4),
            size=72, color=ACCENT_G, font="Georgia")
    txbox(sl, "ACTIVE MEMBERS",
          Inches(0.5), Inches(3.1), Inches(2.5), Inches(0.4),
          size=9, bold=True, color=FG_DARK)

    # Vertical divider
    rect(sl, Inches(3.2), Inches(1.9), Inches(0.04), Inches(1.8), ACCENT_GOLD)

    # Member archetypes
    label_text(sl, "Member Archetypes", Inches(3.5), Inches(1.85), color=ACCENT_GOLD, size=9)

    archetypes = [
        ("🚀", "Founders & Startups"), ("🏥", "Healthcare Providers"),
        ("🧠", "AI Engineers"), ("💼", "Investors & VCs"),
        ("🎓", "Researchers"), ("🏛️", "Policy & Gov"),
    ]
    card_w = Inches(2.3)
    card_h = Inches(0.7)
    cols = 2
    for i, (icon, name) in enumerate(archetypes):
        col = i % cols
        row = i // cols
        cx = Inches(3.5) + col * (card_w + Inches(0.2))
        cy = Inches(2.2) + row * (card_h + Inches(0.15))
        rounded_rect(sl, cx, cy, card_w, card_h, WHITE, line_rgb=BORDER)
        txbox(sl, icon + "  " + name,
              cx + Inches(0.15), cy + Inches(0.12),
              card_w - Inches(0.25), Inches(0.5),
              size=12, color=FG_DARK, bold=False)

    # Geo tags
    rect(sl, Inches(0.5), Inches(6.5), W - Inches(1.0), Inches(0.03),
         RGBColor(0xe0, 0xda, 0xcc))
    txbox(sl, "Global footprint:  SA · UAE · OM · IN · EU · USA",
          Inches(0.5), Inches(6.6), Inches(12), Inches(0.7),
          size=11, color=FG_MUTED)


# ══════════════════════════════════════════════════════
#  S5 — HOW WE WORK
# ══════════════════════════════════════════════════════
def slide_how_we_work(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_CREAM)

    label_text(sl, "How We Work", Inches(0.5), Inches(0.45), color=ACCENT_GOLD, size=9)
    heading(sl, "From Connection to Impact",
            Inches(0.5), Inches(0.7), W - Inches(1.0), Inches(1.0),
            size=36, color=FG_DARK, font="Georgia")

    steps = [
        ("🔗", "Connect", ["Vetted members", "Founders ↔ investors", "Six markets, one room"], True),
        ("🤝", "Collaborate", ["Joint ventures", "Clinical partnerships", "Shared research"], False),
        ("📈", "Grow", ["Funding access", "Talent pipelines", "New markets"], False),
        ("🌍", "Impact", ["Faster adoption", "Vision 2030", "Better patient care"], True),
    ]

    card_w = Inches(2.75)
    card_h = Inches(4.5)
    gap    = Inches(0.4)
    start_x = Inches(0.5)

    for i, (icon, title, items, dark) in enumerate(steps):
        x = start_x + i * (card_w + gap)
        bg_c = BG_DARK if dark else WHITE
        rounded_rect(sl, x, Inches(1.85), card_w, card_h, bg_c,
                     line_rgb=None if dark else BORDER)

        # Icon
        txbox(sl, icon, x + Inches(0.2), Inches(2.05), Inches(1), Inches(0.7), size=28)

        # Title
        title_c = ACCENT_GOLDF if dark else ACCENT_G
        txbox(sl, title, x + Inches(0.2), Inches(2.75), card_w - Inches(0.35), Inches(0.55),
              size=18, bold=True, color=title_c, font="Georgia")

        # Items
        for j, item in enumerate(items):
            item_c = RGBColor(0x9a, 0xb5, 0x9a) if dark else FG_MUTED
            txbox(sl, "• " + item,
                  x + Inches(0.2), Inches(3.35) + j * Inches(0.58),
                  card_w - Inches(0.35), Inches(0.55),
                  size=11, color=item_c)

        # Arrow between cards
        if i < len(steps) - 1:
            arr_x = x + card_w + Inches(0.05)
            txbox(sl, "→", arr_x, Inches(3.7), Inches(0.3), Inches(0.5),
                  size=18, color=ACCENT_GOLD, align=PP_ALIGN.CENTER)

    txbox(sl,
          "Turning a fragmented ecosystem into a connected, intelligent network",
          Inches(0.5), Inches(6.55), W - Inches(1.0), Inches(0.7),
          size=12, color=FG_MUTED, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════
#  S-WEB — WEBINAR PROGRAMME
# ══════════════════════════════════════════════════════
def slide_webinar(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_DARK)

    label_text(sl, "Knowledge Series", Inches(0.5), Inches(0.45), color=ACCENT_GOLDF, size=9)
    heading(sl, "Webinar Programme",
            Inches(0.5), Inches(0.75), Inches(10), Inches(1.0),
            size=36, font="Georgia")

    # KPI cards
    kpis = [
        ("15", "Total Sessions",  RGBColor(0x16, 0x2d, 0x20), WHITE),
        ("10", "Completed",       RGBColor(0x16, 0x2d, 0x20), ACCENT_GOLD),
        ("5",  "Upcoming → Oct 2025", RGBColor(0x16, 0x2d, 0x20), RGBColor(0x8a, 0xb8, 0x9a)),
    ]
    kpi_w = Inches(3.8)
    kpi_h = Inches(1.8)
    for i, (num, lbl, bg, nc) in enumerate(kpis):
        kx = Inches(0.5) + i * (kpi_w + Inches(0.2))
        rounded_rect(sl, kx, Inches(1.75), kpi_w, kpi_h, bg,
                     line_rgb=RGBColor(0x2e, 0x50, 0x3a))
        heading(sl, num, kx + Inches(0.25), Inches(1.82), kpi_w - Inches(0.4), Inches(1.1),
                size=52, color=nc, font="Georgia")
        txbox(sl, lbl.upper(),
              kx + Inches(0.25), Inches(3.0), kpi_w - Inches(0.3), Inches(0.4),
              size=9, bold=True, color=RGBColor(0x88, 0xa0, 0x88))

    # Session dot track legend
    done_dot = "●" * 10
    upco_dot = "○" * 5
    txbox(sl, done_dot + "  " + upco_dot + "   10 / 15 sessions complete",
          Inches(0.5), Inches(3.55), Inches(10), Inches(0.45),
          size=11, color=RGBColor(0x99, 0x99, 0x88))

    # Topic category cards
    cats = [
        ("🤖", "Agentic AI",            "Agentic use in healthcare · How clinicians build AI"),
        ("⚖️", "Regulation",            "Regulatory pathways · Accountability frameworks"),
        ("🚀", "Innovation to Adoption","Why AI stalls after pilot · Scaling AI to patients (Novartis)"),
        ("🔭", "Future Trends",         "What's next for healthcare AI"),
        ("🏥", "Specialty AI",          "Dentistry · Nursing · DATASure"),
    ]
    cat_w = Inches(2.3)
    cat_h = Inches(2.35)
    cat_y = Inches(4.05)
    for i, (icon, name, sub) in enumerate(cats):
        cx = Inches(0.5) + i * (cat_w + Inches(0.18))
        rounded_rect(sl, cx, cat_y, cat_w, cat_h,
                     RGBColor(0x0d, 0x28, 0x1b),
                     line_rgb=RGBColor(0x1e, 0x40, 0x2b))
        txbox(sl, icon, cx + Inches(0.18), cat_y + Inches(0.18),
              Inches(0.6), Inches(0.6), size=22)
        txbox(sl, name, cx + Inches(0.18), cat_y + Inches(0.75),
              cat_w - Inches(0.28), Inches(0.5),
              size=12, bold=True, color=WHITE)
        txbox(sl, sub, cx + Inches(0.18), cat_y + Inches(1.2),
              cat_w - Inches(0.28), Inches(1.0),
              size=9, color=RGBColor(0x7a, 0x92, 0x7a))


# ══════════════════════════════════════════════════════
#  S6 — IMPACT / STATS
# ══════════════════════════════════════════════════════
def slide_impact(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, RGBColor(0x07, 0x1c, 0x11))

    # Header row
    label_text(sl, "Our Impact", Inches(0.5), Inches(0.55), color=ACCENT_GOLD, size=9)
    rect(sl, Inches(1.8), Inches(0.72), Inches(8.5), Inches(0.04), RGBColor(0x2a, 0x4a, 0x30))
    txbox(sl, "Community by the Numbers",
          Inches(10.5), Inches(0.52), Inches(2.7), Inches(0.45),
          size=13, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)

    kpis = [
        ("300+", "Active Members",
         "Clinicians, builders, founders, and investors — vetted and active",
         ACCENT_GOLD, "↑ Growing"),
        ("6",    "Countries & Markets",
         "SA · UAE · Oman · India · Europe · USA — with more joining",
         ACCENT_G, "↑ Expanding"),
        ("4+",   "Member Archetypes",
         "Founders · Clinicians · AI Engineers · Investors — all in one room",
         ACCENT_GOLDF, "↑ Connecting"),
    ]

    kpi_w = Inches(3.8)
    kpi_h = Inches(5.3)
    for i, (num, lbl, desc, ac, trend) in enumerate(kpis):
        kx = Inches(0.5) + i * (kpi_w + Inches(0.24))
        ky = Inches(1.05)
        rounded_rect(sl, kx, ky, kpi_w, kpi_h,
                     RGBColor(0x0d, 0x27, 0x18),
                     line_rgb=RGBColor(0x1e, 0x40, 0x28))
        # accent bar top
        rect(sl, kx, ky, kpi_w, Inches(0.06), ac)

        # Trend badge
        rounded_rect(sl, kx + Inches(0.2), ky + Inches(0.22), Inches(1.4), Inches(0.38),
                     RGBColor(0x0a, 0x38, 0x1e))
        txbox(sl, trend, kx + Inches(0.22), ky + Inches(0.18), Inches(1.35), Inches(0.4),
              size=9, bold=True, color=RGBColor(0x4a, 0xde, 0x80))

        # Big number
        heading(sl, num, kx + Inches(0.2), ky + Inches(0.72),
                kpi_w - Inches(0.35), Inches(1.65),
                size=62, color=WHITE, font="Georgia")

        txbox(sl, lbl, kx + Inches(0.2), ky + Inches(2.35),
              kpi_w - Inches(0.35), Inches(0.5),
              size=14, bold=True, color=RGBColor(0xdd, 0xdd, 0xdd))
        txbox(sl, desc, kx + Inches(0.2), ky + Inches(2.85),
              kpi_w - Inches(0.35), Inches(1.2),
              size=10, color=RGBColor(0x77, 0x88, 0x77))

    txbox(sl, "Community data as of 2025 · GCC Health AI Community Profile",
          Inches(0.5), Inches(6.7), Inches(10), Inches(0.5),
          size=9, color=RGBColor(0x44, 0x55, 0x44))


# ══════════════════════════════════════════════════════
#  S7 — ROADMAP (timeline)
# ══════════════════════════════════════════════════════
def slide_roadmap(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_CREAM)

    heading(sl, "Roadmap",
            Inches(0.5), Inches(0.45), Inches(5), Inches(0.9),
            size=38, color=FG_DARK, font="Georgia")
    txbox(sl, "2023 → 2027+",
          Inches(10), Inches(0.55), Inches(3.1), Inches(0.5),
          size=14, color=FG_MUTED, align=PP_ALIGN.RIGHT)

    txbox(sl,
          "Our mission: make the GCC a global leader in AI-driven healthcare — "
          "the connective tissue linking talent, capital, and clinical institutions across the region.",
          Inches(0.5), Inches(1.45), Inches(12.3), Inches(0.7),
          size=12, color=FG_MUTED)

    # Gradient rule
    rect(sl, Inches(0.5), Inches(2.25), Inches(12.3), Inches(0.04), ACCENT_GOLD)

    # Timeline steps
    steps = [
        ("done",    "2023–2024", "Foundation", "Community launch · 300+ members · 6-market presence", "✓"),
        ("done",    "2024–2025", "Platform",   "Digital hub · Matchmaking tools · Events & summits", "✓"),
        ("current", "2025–2026", "Ecosystem",  "Deal flow pipeline · Research partnerships · 1,000+ members", "3"),
        ("future",  "2027+",     "Leadership", "GCC Health AI index · Policy influence · Global benchmark", "4"),
    ]

    step_w = Inches(3.0)
    step_x_start = Inches(0.5)
    gap = Inches(0.2)
    node_y = Inches(2.6)

    # Connecting line
    rect(sl, Inches(0.9), node_y + Inches(0.3), Inches(11.4), Inches(0.06),
         RGBColor(0xdd, 0xd7, 0xca))
    # Progress line (half)
    rect(sl, Inches(0.9), node_y + Inches(0.3), Inches(5.7), Inches(0.06),
         ACCENT_G)

    for i, (state, period, title, desc, node_lbl) in enumerate(steps):
        sx = step_x_start + i * (step_w + gap)

        # Node circle
        if state == "done":
            circ_c = ACCENT_G
            txt_c  = WHITE
        elif state == "current":
            circ_c = WHITE
            txt_c  = ACCENT_G
        else:
            circ_c = RGBColor(0xe5, 0xe0, 0xd8)
            txt_c  = FG_MUTED

        node_cx = sx + step_w / 2 - Inches(0.32)
        circ = rounded_rect(sl, node_cx, node_y + Inches(0.02),
                            Inches(0.65), Inches(0.65), circ_c)
        txbox(sl, node_lbl, node_cx + Inches(0.01), node_y,
              Inches(0.63), Inches(0.65),
              size=14, bold=True, color=txt_c, align=PP_ALIGN.CENTER)

        # Period
        period_c = ACCENT_G if state in ("done", "current") else FG_MUTED
        txbox(sl, period,
              sx, node_y + Inches(0.82), step_w, Inches(0.45),
              size=9, bold=True, color=period_c, align=PP_ALIGN.CENTER)

        # Title
        title_c = FG_DARK if state != "future" else FG_MUTED
        txbox(sl, title, sx, node_y + Inches(1.25), step_w, Inches(0.55),
              size=18, bold=True, color=title_c, align=PP_ALIGN.CENTER,
              font="Georgia")

        # Desc
        txbox(sl, desc, sx + Inches(0.1), node_y + Inches(1.8), step_w - Inches(0.2), Inches(1.4),
              size=10, color=FG_MUTED, align=PP_ALIGN.CENTER)

        # In Progress badge
        if state == "current":
            badge = rounded_rect(sl, sx + step_w/2 - Inches(0.75), node_y + Inches(3.3),
                                 Inches(1.5), Inches(0.38), ACCENT_G)
            txbox(sl, "IN PROGRESS",
                  sx + step_w/2 - Inches(0.72), node_y + Inches(3.28),
                  Inches(1.44), Inches(0.38),
                  size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Footer progress
    rect(sl, Inches(0.5), Inches(6.75), Inches(12.3), Inches(0.05),
         RGBColor(0xe8, 0xe1, 0xd2))
    txbox(sl, "Overall Progress",
          Inches(0.5), Inches(6.87), Inches(2.5), Inches(0.45),
          size=10, bold=True, color=ACCENT_G)
    # Progress bar track
    rect(sl, Inches(3.2), Inches(6.98), Inches(8.8), Inches(0.25),
         RGBColor(0xe8, 0xe1, 0xd2))
    rect(sl, Inches(3.2), Inches(6.98), Inches(4.4), Inches(0.25),
         ACCENT_G)
    txbox(sl, "50%",
          Inches(12.1), Inches(6.87), Inches(1.1), Inches(0.45),
          size=10, bold=True, color=FG_DARK, align=PP_ALIGN.RIGHT)


# ══════════════════════════════════════════════════════
#  S8 — TEAM
# ══════════════════════════════════════════════════════
def slide_team(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_CREAM)

    label_text(sl, "The Team", Inches(0.5), Inches(0.45), color=ACCENT_GOLD, size=9)
    heading(sl, "Leadership",
            Inches(0.5), Inches(0.7), Inches(8), Inches(0.9),
            size=38, color=FG_DARK, font="Georgia")

    txbox(sl, "The team turning a regional vision into a working network",
          Inches(0.5), Inches(1.65), Inches(12), Inches(0.5),
          size=13, color=FG_MUTED, align=PP_ALIGN.CENTER)

    # Gold rule
    rect(sl, Inches(2), Inches(2.25), Inches(9.3), Inches(0.05), ACCENT_GOLD)

    team = [
        ("ZH", "Dr. Zahid Hussain",   "Founder & Community Lead",
         "Sets the clinical vision and rallies the network behind it"),
        ("MA", "Mohammed Almosa",     "Operations & Growth Lead",
         "Builds the partnerships and systems that scale the community"),
        ("ME", "Maria Expósito",      "Strategy & Finance Lead",
         "Shapes strategy and keeps the community's growth sustainable"),
    ]

    card_w = Inches(3.7)
    card_h = Inches(4.3)
    start_x = Inches(0.6)
    gap = Inches(0.4)
    card_y = Inches(2.45)

    for i, (initials, name, role, bio) in enumerate(team):
        cx = start_x + i * (card_w + gap)
        rounded_rect(sl, cx, card_y, card_w, card_h, WHITE, line_rgb=BORDER)

        # Avatar circle
        av_size = Inches(1.3)
        av_x = cx + card_w / 2 - av_size / 2
        av_y = card_y + Inches(0.35)
        circ = rounded_rect(sl, av_x, av_y, av_size, av_size,
                            RGBColor(0xe8, 0xf5, 0xec))
        txbox(sl, initials,
              av_x, av_y - Inches(0.05), av_size, av_size + Inches(0.05),
              size=22, bold=True, color=ACCENT_G, align=PP_ALIGN.CENTER,
              font="Georgia")

        # Name
        txbox(sl, name,
              cx + Inches(0.15), card_y + Inches(1.85),
              card_w - Inches(0.3), Inches(0.5),
              size=13, bold=True, color=ACCENT_G, align=PP_ALIGN.CENTER)

        # Role
        txbox(sl, role,
              cx + Inches(0.15), card_y + Inches(2.32),
              card_w - Inches(0.3), Inches(0.6),
              size=11, color=FG_DARK, align=PP_ALIGN.CENTER)

        # Bio
        txbox(sl, bio,
              cx + Inches(0.2), card_y + Inches(2.95),
              card_w - Inches(0.4), Inches(1.0),
              size=10, color=FG_MUTED, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════
#  S9 — CTA / CLOSE
# ══════════════════════════════════════════════════════
def slide_cta(prs):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, BG_DARK)
    # Right accent
    rect(sl, Inches(7), 0, Inches(6.333), H, RGBColor(0x0a, 0x5c, 0x3a))

    # Watermark "GCC"
    txbox(sl, "GCC",
          Inches(5), Inches(1.8), Inches(8), Inches(6),
          size=200, bold=True, color=RGBColor(0x0e, 0x28, 0x1a),
          font="Calibri")

    # Logo
    txbox(sl, "GCC  Health AI",
          Inches(0.45), Inches(0.3), Inches(3.5), Inches(0.5),
          size=13, bold=True, color=WHITE)

    # Kicker
    txbox(sl, "— GET INVOLVED",
          Inches(0.45), Inches(3.8),
          Inches(5), Inches(0.4),
          size=10, bold=True, color=ACCENT_GOLDF)

    # H2
    heading(sl, "Shape the Future\nof Healthcare\nin the Gulf",
            Inches(0.45), Inches(4.1), Inches(8.5), Inches(2.9),
            size=42, color=WHITE, font="Georgia")

    # Subtitle
    body_text(sl,
        "Join a curated network of the region's brightest minds in healthcare and AI. "
        "Connect, collaborate, and help build what comes next.",
        Inches(0.45), Inches(6.65), Inches(7.5), Inches(0.7),
        size=13, color=RGBColor(0xb8, 0xb8, 0xb8))

    # Contact icons as text
    txbox(sl, "✉  linkedin.com/in/gcc-health-ai  ·  🌐 gcchealth.ai",
          Inches(0.45), Inches(7.1), Inches(8), Inches(0.4),
          size=11, color=ACCENT_GOLDF)

    # Brand
    txbox(sl, "GCC Health AI · 2025",
          Inches(10), Inches(7.1), Inches(3.1), Inches(0.35),
          size=9, color=RGBColor(0x55, 0x66, 0x55),
          align=PP_ALIGN.RIGHT)


# ── Build all slides ───────────────────────────────────
slide_hero(prs)
slide_about(prs)
slide_problem(prs)
slide_members(prs)
slide_how_we_work(prs)
slide_webinar(prs)
slide_impact(prs)
slide_roadmap(prs)
slide_team(prs)
slide_cta(prs)

out = "/home/user/slides/gcc-health-ai-profile.pptx"
prs.save(out)
print(f"Saved: {out}  ({len(prs.slides)} slides)")
