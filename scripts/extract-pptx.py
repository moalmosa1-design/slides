#!/usr/bin/env python3
"""
extract-pptx.py — Extract slide text content from a .pptx file.

Usage:
    python extract-pptx.py <path-to-file.pptx>

Output:
    Prints JSON to stdout in the format:
    {
      "slides": [
        {
          "index": 0,
          "title": "Slide title text",
          "body": ["bullet point 1", "bullet point 2", ...]
        },
        ...
      ]
    }

Dependencies:
    pip install python-pptx

Notes:
    - The "title" is taken from the shape with the placeholder type TITLE or
      CENTER_TITLE. If no title placeholder exists, the first text shape is used.
    - "body" contains text from all non-title text shapes, split by paragraph.
      Empty paragraphs are skipped.
    - Run-level text within a paragraph is concatenated with a single space.
"""

import sys
import json
import os

def extract_pptx(filepath: str) -> dict:
    """
    Parse a .pptx file and return a dict with a 'slides' list.
    Each slide entry has: index (int), title (str), body (list[str]).
    """
    try:
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt
        import pptx.shapes.placeholder as ph_module
        from pptx.enum.shapes import PP_PLACEHOLDER
    except ImportError:
        print(
            json.dumps({
                "error": "python-pptx is not installed. Run: pip install python-pptx"
            }),
            file=sys.stderr
        )
        sys.exit(2)

    if not os.path.isfile(filepath):
        print(
            json.dumps({"error": f"File not found: {filepath}"}),
            file=sys.stderr
        )
        sys.exit(2)

    if not filepath.lower().endswith(".pptx"):
        print(
            json.dumps({"error": f"File does not appear to be a .pptx: {filepath}"}),
            file=sys.stderr
        )
        sys.exit(2)

    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(
            json.dumps({"error": f"Failed to open presentation: {str(e)}"}),
            file=sys.stderr
        )
        sys.exit(1)

    TITLE_TYPES = {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.VERTICAL_TITLE,
    }

    slides_data = []

    for slide_index, slide in enumerate(prs.slides):
        title_text = ""
        body_lines = []

        # Collect all shapes that have text frames
        title_shape = None
        body_shapes = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            # Identify title placeholder
            is_title = False
            if shape.is_placeholder:
                try:
                    ph_type = shape.placeholder_format.type
                    if ph_type in TITLE_TYPES:
                        is_title = True
                except Exception:
                    pass

            if is_title:
                if title_shape is None:  # use first title found
                    title_shape = shape
            else:
                body_shapes.append(shape)

        # Extract title text
        if title_shape is not None:
            title_text = _shape_to_text(title_shape)
        elif body_shapes:
            # Fallback: use first text shape as title
            title_text = _shape_to_text(body_shapes.pop(0))

        # Extract body text — one entry per non-empty paragraph
        for shape in body_shapes:
            for para in shape.text_frame.paragraphs:
                line = _paragraph_to_text(para).strip()
                if line:
                    body_lines.append(line)

        slides_data.append({
            "index": slide_index,
            "title": title_text.strip(),
            "body": body_lines,
        })

    return {"slides": slides_data}


def _shape_to_text(shape) -> str:
    """Concatenate all paragraphs in a shape's text frame into a single string."""
    lines = []
    for para in shape.text_frame.paragraphs:
        text = _paragraph_to_text(para).strip()
        if text:
            lines.append(text)
    return " ".join(lines)


def _paragraph_to_text(paragraph) -> str:
    """Concatenate all runs in a paragraph into a single string."""
    parts = []
    for run in paragraph.runs:
        t = run.text
        if t:
            parts.append(t)
    # Also capture text in paragraph.text (catches cases with no runs)
    if not parts and paragraph.text:
        return paragraph.text
    return " ".join(parts)


def main():
    if len(sys.argv) < 2 or sys.argv[1] in ("-h", "--help"):
        print(
            "Usage: python extract-pptx.py <path-to-file.pptx>\n"
            "Extracts slide text and prints JSON to stdout.\n"
            "\n"
            "Output format:\n"
            '  {"slides": [{"index": 0, "title": "...", "body": ["...", ...]}, ...]}\n'
            "\n"
            "Requirements: pip install python-pptx",
            file=sys.stderr
        )
        sys.exit(0 if "--help" in sys.argv or "-h" in sys.argv else 1)

    filepath = sys.argv[1]
    result = extract_pptx(filepath)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
